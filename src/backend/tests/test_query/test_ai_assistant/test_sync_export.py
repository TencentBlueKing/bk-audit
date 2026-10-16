# -*- coding: utf-8 -*-
"""
TencentBlueKing is pleased to support the open source community by making
蓝鲸智云 - 审计中心 (BlueKing - Audit Center) available.
Copyright (C) 2023 THL A29 Limited,
a Tencent company. All rights reserved.
Licensed under the MIT License (the "License");
you may not use this file except in compliance with the License.
You may obtain a copy of the License at http://opensource.org/licenses/MIT
Unless required by applicable law or agreed to in writing,
software distributed under the License is distributed on an "AS IS" BASIS,
WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
See the License for the specific language governing permissions and limitations
under the License.
We undertake not to change the open source license (MIT license) that applies
to the current version of the project delivered to anyone in the future.

F4.5 全量导出同步直出（SyncExportRunner）组件测试。

零 DB（SimpleTestCase 纪律）+ FakeClock 注入：守卫与降级逻辑全部以假时钟推进验证，
CI 内不出现任何真实等待；Doris/权限全 mock，并发槽走 locmem cache 真实占位/释放。
用例编号对齐设计文档测试矩阵（A1/A3~A9/A11~A13/A16~A19）。
"""

import io
from unittest import mock

import openpyxl
from django.core.cache import cache
from django.test import override_settings

from services.web.query.ai_assistant.constants import (
    AI_ASSISTANT_EXPORT_SOURCE,
    SNAPSHOT_DEFAULT_COLUMNS,
    SYNC_EXPORT_SLOT_KEY_PREFIX,
    SyncExportFallbackReason,
)
from services.web.query.ai_assistant.exceptions import (
    AIOutputInvalidError,
    AIPermissionDeniedError,
    SyncExportFallback,
)
from services.web.query.ai_assistant.services.export import (
    PreviewExportFile,
    SyncExportRunner,
)
from tests.test_query.test_ai_assistant.base import AIAssistantTestCase

EXPORT_MODULE = "services.web.query.ai_assistant.services.export"

# 守卫用例的小预算基线：deadline 10s / 硬超时 20s / count 限时 3s /
# 条数上限 1000 / 分页 100 / 并发 2 / 行均成本 1.5ms（小到不触发首批外推）
SYNC_TEST_SETTINGS = dict(
    AI_EXPORT_SYNC_ENABLED=True,
    AI_EXPORT_SYNC_MAX_COUNT=1000,
    AI_EXPORT_SYNC_DEADLINE_SECONDS=10,
    AI_EXPORT_SYNC_HARD_LIMIT_SECONDS=20,
    AI_EXPORT_SYNC_PAGE_SIZE=100,
    AI_EXPORT_SYNC_CONCURRENCY=2,
    AI_EXPORT_SYNC_ROW_COST_BUDGET_MS=1.5,
    AI_EXPORT_SYNC_COUNT_MAX_SECONDS=3,
)


class FakeClock:
    """假时钟：每次读取推进 tick（模拟单步耗时），支持 advance 快进模拟慢批次/慢查询。"""

    def __init__(self, tick: float = 0.01):
        self._now = 0.0
        self._tick = tick

    def __call__(self) -> float:
        current = self._now
        self._now += self._tick
        return current

    def advance(self, seconds: float):
        self._now += seconds


def make_runner_export_config() -> dict:
    """Runner 输入的最终形态 export_config（MessageExportService._build_export_config 之后的产物）。"""

    return {
        "field_scope": "specified",
        "fields": [
            {"raw_name": raw_name, "display_name": display_name, "keys": []}
            for raw_name, display_name in SNAPSHOT_DEFAULT_COLUMNS
        ],
        "source": AI_ASSISTANT_EXPORT_SOURCE,
    }


def make_rows(count: int) -> list:
    """构造一批 Doris 行数据（缺键由 DataProcessor 兜底空值）。"""

    return [
        {"start_time": "2026-08-24 10:00:00", "username": "admin", "system_id": "bk_log", "result_code": 0}
        for _ in range(count)
    ]


# 类级统一放行权限（守卫用例聚焦时间预算；权限拒绝单独覆盖）
@override_settings(**SYNC_TEST_SETTINGS)
@mock.patch(f"{EXPORT_MODULE}.SearchLogPermission.has_system_search_permission", return_value=True)
@mock.patch(f"{EXPORT_MODULE}.DataFetcher.get_total")
class TestSyncExportRunnerGuard(AIAssistantTestCase):
    """守卫与降级：任何降级信号都以 SyncExportFallback 抛出且附带低基数 reason。"""

    def setUp(self):
        super().setUp()
        # 并发槽 locmem 状态隔离：每个用例从干净槽位集合开始
        cache.clear()

    def make_runner(self, **overrides) -> SyncExportRunner:
        return SyncExportRunner(
            condition=self.make_condition(),
            namespace=self.namespace,
            export_config=overrides.pop("export_config", make_runner_export_config()),
            username=self.username,
        )

    def _assert_fallback(self, runner, clock, reason: str):
        with self.assertRaises(SyncExportFallback) as ctx:
            runner.run(clock=clock)
        self.assertEqual(ctx.exception.reason, reason)
        return ctx.exception

    # ------------------------------------------------------------------
    # A1/A13/A16：同步直出成功 + 产物内容正确性 + 指标
    # ------------------------------------------------------------------

    @mock.patch(f"{EXPORT_MODULE}.DataFetcher.fetch_patch_logs")
    @mock.patch(f"{EXPORT_MODULE}._report_sync_export_metric")
    def test_sync_export_success_returns_file(self, mock_metric, mock_fetch, mock_total, mock_perm):
        """A1：预算内两批数据直出 XLSX；行数/表头/单行 AI 样式/指标全断言。"""

        mock_total.return_value = 200
        mock_fetch.side_effect = lambda query_params, page, page_size: make_rows(page_size)

        result = self.make_runner().run(clock=FakeClock())

        self.assertIsInstance(result, PreviewExportFile)
        self.assertTrue(result.file_name.endswith(".xlsx"))
        # 数据行正确性：表头 1 行 + 200 数据行
        workbook = openpyxl.load_workbook(io.BytesIO(result.content))
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 1 + 200)
        titles = [cell.value for cell in sheet[1]]
        # A13：AI 单行表头（中文显示名(英文字段路径)），与异步导出同构
        self.assertIn("操作起始时间(start_time)", titles)
        self.assertIn("操作人(username)", titles)
        first_row = [cell.value for cell in sheet[2]]
        self.assertIn("admin", first_row)
        # A16：成功指标（rows 携带导出条数）
        mock_metric.assert_called_once()
        self.assertEqual(mock_metric.call_args.kwargs["status"], "success")
        self.assertEqual(mock_metric.call_args.kwargs["rows"], 200)
        self.assertEqual(mock_fetch.call_count, 2)

    @mock.patch(f"{EXPORT_MODULE}.DataFetcher.fetch_patch_logs")
    def test_sync_export_zero_total_returns_empty_file(self, mock_fetch, mock_total, mock_perm):
        """零数据快照沿用正常直出（对齐预览导出零数据语义）：仅表头、无数据行。"""

        mock_total.return_value = 0

        result = self.make_runner().run(clock=FakeClock())

        workbook = openpyxl.load_workbook(io.BytesIO(result.content))
        self.assertEqual(workbook.active.max_row, 1)
        mock_fetch.assert_not_called()

    # ------------------------------------------------------------------
    # 准入与 count 守卫
    # ------------------------------------------------------------------

    @mock.patch(f"{EXPORT_MODULE}.DataFetcher.fetch_patch_logs")
    @mock.patch(f"{EXPORT_MODULE}._report_sync_export_metric")
    def test_count_total_exceed_falls_back(self, mock_metric, mock_fetch, mock_total, mock_perm):
        """A3：count 复检超条数上限（快照 total 漂移以导出时点为准）→ 降级且不拉数据。"""

        mock_total.return_value = 101

        with override_settings(AI_EXPORT_SYNC_MAX_COUNT=100):
            self._assert_fallback(self.make_runner(), FakeClock(), SyncExportFallbackReason.TOTAL_EXCEED)
        mock_fetch.assert_not_called()
        self.assertEqual(mock_metric.call_args.kwargs["reason"], SyncExportFallbackReason.TOTAL_EXCEED)

    @mock.patch(f"{EXPORT_MODULE}.DataFetcher.fetch_patch_logs")
    def test_count_slow_falls_back(self, mock_fetch, mock_total, mock_perm):
        """A4：count 复检耗时超集群慢阈值 → 降级（数据量小但集群慢不同步尝试）。"""

        clock = FakeClock()

        def slow_count(query_params):
            clock.advance(5)
            return 10

        mock_total.side_effect = slow_count

        self._assert_fallback(self.make_runner(), clock, SyncExportFallbackReason.COUNT_SLOW)
        mock_fetch.assert_not_called()

    @mock.patch(f"{EXPORT_MODULE}.DataFetcher.fetch_patch_logs")
    def test_first_batch_estimate_exceeds_deadline(self, mock_fetch, mock_total, mock_perm):
        """A7：首批行均成本外推超预算（数据量与行均成本双超）→ 拉取前拦截。"""

        mock_total.return_value = 500

        with override_settings(AI_EXPORT_SYNC_ROW_COST_BUDGET_MS=100):
            self._assert_fallback(self.make_runner(), FakeClock(), SyncExportFallbackReason.DEADLINE)
        mock_fetch.assert_not_called()

    # ------------------------------------------------------------------
    # 批间守卫（每批开始前：绝对预算 + 上批耗时外推）
    # ------------------------------------------------------------------

    @mock.patch(f"{EXPORT_MODULE}.DataFetcher.fetch_patch_logs")
    def test_inter_batch_absolute_budget_falls_back(self, mock_fetch, mock_total, mock_perm):
        """A18：单批波动大 → 下一批开始前 elapsed 超绝对预算（deadline-收尾预留）→ 立即降级。"""

        clock = FakeClock()
        mock_total.return_value = 300  # 3 批

        def slow_first_batch(query_params, page, page_size):
            if page == 1:
                clock.advance(8.1)
            return make_rows(page_size)

        mock_fetch.side_effect = slow_first_batch

        self._assert_fallback(self.make_runner(), clock, SyncExportFallbackReason.DEADLINE)
        # 第一批已拉完，第二批开始前拦截
        self.assertEqual(mock_fetch.call_count, 1)

    @mock.patch(f"{EXPORT_MODULE}.DataFetcher.fetch_patch_logs")
    def test_inter_batch_extrapolation_falls_back_first(self, mock_fetch, mock_total, mock_perm):
        """A6：绝对预算未触发（elapsed<deadline-2）但上批耗时外推超预算 → 外推先拦。"""

        clock = FakeClock()
        mock_total.return_value = 300  # 3 批

        def steady_slow_batch(query_params, page, page_size):
            clock.advance(4.1)
            return make_rows(page_size)

        mock_fetch.side_effect = steady_slow_batch

        self._assert_fallback(self.make_runner(), clock, SyncExportFallbackReason.DEADLINE)
        # elapsed≈4.2 < 8（绝对预算未触发），4.2 + 4.1*1.5 > 10（外推触发）
        self.assertEqual(mock_fetch.call_count, 1)

    # ------------------------------------------------------------------
    # A17/A19：边界分级——最后一批波动超额完成 vs 收尾超硬上界
    # ------------------------------------------------------------------

    @mock.patch(f"{EXPORT_MODULE}.DataFetcher.fetch_patch_logs")
    def test_last_batch_overrun_completes_within_hard_limit(self, mock_fetch, mock_total, mock_perm):
        """A17：最后一批波动超软超时但数据已拉完 → 只差收尾不降级（硬上界内超额完成）。"""

        clock = FakeClock()
        mock_total.return_value = 200  # 2 批

        def wavy_batch(query_params, page, page_size):
            if page == 1:
                clock.advance(1)
            else:
                clock.advance(9)  # 批中波动：结束时 elapsed≈10 超软超时
            return make_rows(page_size)

        mock_fetch.side_effect = wavy_batch

        result = self.make_runner().run(clock=clock)

        # 此时中止是双输（用户白等 30s 还要重走异步），软超时优雅：能完成就完成
        self.assertIsInstance(result, PreviewExportFile)
        workbook = openpyxl.load_workbook(io.BytesIO(result.content))
        self.assertEqual(workbook.active.max_row, 1 + 200)
        self.assertEqual(mock_fetch.call_count, 2)

    @mock.patch(f"{EXPORT_MODULE}.DataFetcher.fetch_patch_logs")
    def test_finalize_over_hard_limit_falls_back(self, mock_fetch, mock_total, mock_perm):
        """A19：数据拉完但收尾阶段超硬超时绝对天花板 → 果断降级，请求永不挂死。"""

        clock = FakeClock()
        mock_total.return_value = 200  # 2 批

        def heavy_batch(query_params, page, page_size):
            if page == 1:
                clock.advance(1)
            else:
                clock.advance(19)  # 结束 elapsed≈20 超硬超时 20s
            return make_rows(page_size)

        mock_fetch.side_effect = heavy_batch

        self._assert_fallback(self.make_runner(), clock, SyncExportFallbackReason.HARD_LIMIT)

    # ------------------------------------------------------------------
    # A8/A9：全局并发槽（locmem 真实占位/释放）
    # ------------------------------------------------------------------

    @mock.patch(f"{EXPORT_MODULE}.DataFetcher.fetch_patch_logs")
    def test_concurrency_limit_falls_back(self, mock_fetch, mock_total, mock_perm):
        """A8：并发槽满 → 立即降级（不排队不等待），且不查 Doris。"""

        with override_settings(AI_EXPORT_SYNC_CONCURRENCY=1):
            # 预占唯一槽位（模拟另一请求在途）
            self.assertTrue(cache.add(f"{SYNC_EXPORT_SLOT_KEY_PREFIX}:0", "occupied-by-other", timeout=60))
            mock_total.return_value = 10

            self._assert_fallback(self.make_runner(), FakeClock(), SyncExportFallbackReason.CONCURRENCY)
            mock_total.assert_not_called()

    @mock.patch(f"{EXPORT_MODULE}.DataFetcher.fetch_patch_logs")
    def test_concurrency_slot_released_after_success(self, mock_fetch, mock_total, mock_perm):
        """A9：成功路径结束后释放槽位（槽位可被下一请求复用）。"""

        mock_total.return_value = 10
        mock_fetch.side_effect = lambda query_params, page, page_size: make_rows(page_size)

        self.make_runner().run(clock=FakeClock())

        self.assertIsNone(cache.get(f"{SYNC_EXPORT_SLOT_KEY_PREFIX}:0"))
        self.assertIsNone(cache.get(f"{SYNC_EXPORT_SLOT_KEY_PREFIX}:1"))

    @mock.patch(f"{EXPORT_MODULE}.DataFetcher.fetch_patch_logs")
    def test_concurrency_slot_released_after_fallback(self, mock_fetch, mock_total, mock_perm):
        """A9：降级路径同样释放槽位（降级不泄漏槽位）。"""

        mock_total.return_value = 101  # total_exceed 降级

        with override_settings(AI_EXPORT_SYNC_MAX_COUNT=100):
            self._assert_fallback(self.make_runner(), FakeClock(), SyncExportFallbackReason.TOTAL_EXCEED)

        self.assertIsNone(cache.get(f"{SYNC_EXPORT_SLOT_KEY_PREFIX}:0"))
        self.assertIsNone(cache.get(f"{SYNC_EXPORT_SLOT_KEY_PREFIX}:1"))

    # ------------------------------------------------------------------
    # A11/A12：权限与白名单（校验失败报错，不降级）
    # ------------------------------------------------------------------

    @mock.patch(f"{EXPORT_MODULE}.DataFetcher.fetch_patch_logs")
    def test_permission_denied_not_swallowed(self, mock_fetch, mock_total, mock_perm):
        """A11：无权限 → AIPermissionDeniedError 原样传播（不吞异常不降级异步）。"""

        mock_total.return_value = 10
        # 类级默认放行权限；本用例显式改为拒绝（同目标方法级叠加不可靠，直接改类级 mock）
        mock_perm.return_value = False

        with self.assertRaises(AIPermissionDeniedError):
            self.make_runner().run(clock=FakeClock())
        mock_fetch.assert_not_called()

    @mock.patch(f"{EXPORT_MODULE}.DataFetcher.fetch_patch_logs")
    def test_invalid_fields_rejected(self, mock_fetch, mock_total, mock_perm):
        """A12：越权字段 → AIOutputInvalidError 原样传播（与异步路径同源白名单）。"""

        mock_total.return_value = 10
        export_config = make_runner_export_config()
        export_config["fields"] = [{"raw_name": "hack_field", "display_name": "x", "keys": []}]

        with self.assertRaises(AIOutputInvalidError):
            self.make_runner(export_config=export_config).run(clock=FakeClock())
        mock_fetch.assert_not_called()
